from pathlib import Path
from .utils import now_date, slugify, ensure_unique_path, read_text_safe

def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        return "PDF erkannt. Paket fehlt: pip install pypdf"
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        pages.append(f"## Seite {i}\n\n{page.extract_text() or ''}")
    return "\n\n".join(pages).strip() or "Kein Text extrahierbar. OCR erforderlich."

def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception:
        return "DOCX erkannt. Paket fehlt: pip install python-docx"
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except Exception:
        return "XLSX erkannt. Paket fehlt: pip install openpyxl"
    wb = openpyxl.load_workbook(str(path), data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)

def extract_ocr_image(path: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
    except Exception:
        return "OCR benötigt: pip install pillow pytesseract plus Tesseract Installation."
    return pytesseract.image_to_string(Image.open(path), lang="deu+eng")

def import_document_to_inbox(settings: dict, file_path: str) -> Path:
    source = Path(file_path)
    inbox = Path(settings["inbox_path"]) / "Documents"
    inbox.mkdir(parents=True, exist_ok=True)

    suffix = source.suffix.lower()
    if suffix == ".pdf":
        text = extract_pdf(source)
    elif suffix == ".docx":
        text = extract_docx(source)
    elif suffix == ".xlsx":
        text = extract_xlsx(source)
    elif suffix in [".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"]:
        text = extract_ocr_image(source)
    else:
        text = read_text_safe(source)

    target = ensure_unique_path(inbox / f"{now_date()}_{slugify(source.stem)}.md")
    md = f"""---
title: "{source.stem}"
type: document_import
source_file: "{source}"
created: {now_date()}
tags:
  - document
---

# {source.stem}

{text}
"""
    target.write_text(md, encoding="utf-8")
    return target
