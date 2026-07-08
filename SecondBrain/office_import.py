from pathlib import Path
from .utils import read_text_safe

def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception:
        return "DOCX erkannt. Installation erforderlich: pip install python-docx"
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except Exception:
        return "XLSX erkannt. Installation erforderlich: pip install openpyxl"
    wb = openpyxl.load_workbook(str(path), data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)

def extract_office(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".pptx":
        return "PPTX erkannt. Produktiver Import wird vorbereitet."
    return read_text_safe(path)
